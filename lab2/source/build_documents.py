"""Build the Lab 2 deliverables and the verbatim AI recommendation report."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def setup(doc):
    for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
        for border in list(s.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
        for color in s.element.iter(qn('w:color')):
            for attr in ['themeColor','themeTint','themeShade']: color.attrib.pop(qn('w:'+attr),None)
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    for name,size in [('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size)
    sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.7)
    p=sec.footer.paragraphs[0]; p.alignment=2
    p.add_run('PlugPlan SG | Lab 2 | ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)
    for r in p.runs: r.font.size=Pt(8)
    doc.core_properties.author='SC2006 Team SCE3-04'

def p(doc,text): doc.add_paragraph(text)

def page_map():
    path=ROOT/'source/page_map.json'
    if not path.exists(): return {}
    try:
        data=json.loads(path.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}
    if isinstance(data,dict):
        data=data.get('deliverables',data.get('page_ranges',data))
    if isinstance(data,dict): return data
    if isinstance(data,list):
        return {item.get('key'):item.get('page_range',item.get('pages','TBD'))
                for item in data if isinstance(item,dict) and item.get('key')}
    return {}

def page_range(mapping,key):
    value=mapping.get(key,'TBD')
    if isinstance(value,dict): value=value.get('page_range',value.get('pages','TBD'))
    if isinstance(value,(list,tuple)) and len(value)==2: return f'{value[0]}–{value[1]}'
    return str(value)

def model_section(text,number):
    match=re.search(rf'^##\s+{number}\.\s+[^\n]*\n(.*?)(?=^##\s+\d+\.\s|\Z)',text,re.M|re.S)
    return match.group(1).strip() if match else ''

def display_title(item):
    """Keep generated diagram headings aligned with the four document sections."""
    section=item['section']; raw=item.get('title','')
    if section=='usecase': return '1 Use Case — Complete Use Case Diagram'
    if section=='entity':
        suffix=re.sub(r'^2\s+(?:Conceptual(?: Entity)? Model|Class diagram of entity classes)\s*', '', raw, flags=re.I)
        return '2 Class diagram of entity classes'+(f' — {suffix}' if suffix else '')
    if section=='boundary': return '3 Key boundary and control classes'
    if section=='dialog':
        suffix=re.sub(r'^4\s+Initial Dialog Map\s*', '', raw, flags=re.I)
        return '4 Initial Dialog Map'+(f' — {suffix}' if suffix else '')
    return raw

def bullets(doc,items):
    for item in items: doc.add_paragraph(item,style='List Bullet')
def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    if widths:
        t.autofit=False
        for col,width in zip(t.columns,widths): col.width=Inches(width)
    for i,h in enumerate(headers): t.rows[0].cells[i].text=h
    for row in rows:
        for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[ci])
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF2'); pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4); para.paragraph_format.space_before=Pt(4)
                for run in para.runs: run.font.size=Pt(9); run.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def md(doc,text):
    lines=text.splitlines(); i=0
    def clean(text):
        text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text)
        return text.replace('`','').replace('**','')
    while i<len(lines):
        line=lines[i]; i+=1
        if not line.strip(): continue
        if line.startswith('|'):
            chunk=[line]
            while i<len(lines) and lines[i].startswith('|'): chunk.append(lines[i]); i+=1
            cells=lambda row:[clean(c.strip()) for c in re.split(r'(?<!\\)\|',row.strip().strip('|'))]
            headers=cells(chunk[0]); widths=None
            if len(headers)==2: widths=[2.0,4.8] if 'Association' not in headers[0] else [3.0,3.8]
            elif len(headers)==3: widths=[1.85,3.3,1.65]
            elif len(headers)==4: widths=[1.7,1.8,2.3,1.0]
            if headers[0]=='From': widths=[.6,3.15,3.05]
            if headers[0]=='Use case': widths=[1.85,1.8,3.15]
            if headers[0]=='State': widths=[2.0,1.0,3.8]
            table(doc,headers,[cells(r) for r in chunk[2:]],widths)
            continue
        line=clean(line)
        if line.startswith('#'):
            level=min(len(line)-len(line.lstrip('#')),3); doc.add_heading(line.lstrip('# ').strip(),level)
        elif line.startswith('- '):
            while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','- ','|')):
                line+=' '+clean(lines[i]); i+=1
            doc.add_paragraph(line[2:],style='List Bullet')
        else:
            while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','- ','|')):
                line+=' '+clean(lines[i]); i+=1
            p(doc,line)

def diagram(doc,path,title,caption):
    from PIL import Image
    sec=doc.add_section(WD_SECTION.NEW_PAGE); sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
    sec.left_margin=sec.right_margin=Inches(.55); sec.top_margin=sec.bottom_margin=Inches(.5)
    doc.add_heading(title,1)
    im=Image.open(path); w,h=im.size; maxw,maxh=10.45,5.65
    width=min(maxw,maxh*w/h)
    para=doc.add_paragraph(); para.alignment=1; para.add_run().add_picture(str(path),width=Inches(width))
    p(doc,caption)

def portrait(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE); sec.orientation=WD_ORIENT.PORTRAIT
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.left_margin=sec.right_margin=Inches(.7); sec.top_margin=sec.bottom_margin=Inches(.65)

def build_srs():
    doc=Document(); setup(doc)
    doc.add_heading('Deliverables',0)
    p(doc,'SC2006 Team SCE3-04')
    mapping=page_map()
    table(doc,['Deliverable','Page range'],[
        ('Complete Use Case diagram',page_range(mapping,'usecase-diagram')),
        ('Use Case descriptions (UC-01–UC-08)',page_range(mapping,'usecases')),
        ('Class diagram of entity classes',page_range(mapping,'entity')),
        ('Key boundary and control classes',page_range(mapping,'boundary')),
        ('Initial Dialog Map',page_range(mapping,'dialog'))],[4.9,1.8])
    manifest=json.loads((ROOT/'diagrams/manifest.json').read_text(encoding='utf-8')) if (ROOT/'diagrams/manifest.json').exists() else []
    for item in manifest:
        if item['section']=='usecase': diagram(doc,ROOT/item['file'],display_title(item),item['caption'])
    portrait(doc)
    for index,u in enumerate(json.loads((ROOT/'source/use_cases.json').read_text(encoding='utf-8'))):
        if index: doc.add_page_break()
        doc.add_heading(u['id']+' '+u['name'],1)
        p(doc,u['description'])
        table(doc,['Field','Specification'],[('Primary and supporting actors',u['actor']),('Priority and frequency',u['priority']+'; '+u['frequency']),('Traceability',u['trace'])],[1.65,5.15])
        for title,key in [('Preconditions','pre'),('Postconditions','post')]: doc.add_heading(title,2); bullets(doc,u[key])
        doc.add_heading('Normal flow of events',2)
        for i,text in enumerate(u['flow'],1): p(doc,f'{i}. {text}')
        for title,key in [('Alternative flows','alts'),('Exceptions','exceptions')]: doc.add_heading(title,2); bullets(doc,u[key])
        doc.add_heading('Special requirements and assumptions',2); p(doc,u['special']); p(doc,u['assumptions'])
    notes=(ROOT/'source/model_spec.md').read_text(encoding='utf-8') if (ROOT/'source/model_spec.md').exists() else ''
    for section in ('entity','boundary','dialog'):
        for item in manifest:
            if item['section']==section: diagram(doc,ROOT/item['file'],display_title(item),item['caption'])
        portrait(doc)
        doc.add_heading({'entity':'2 Class diagram of entity classes','boundary':'3 Key boundary and control classes','dialog':'4 Initial Dialog Map'}[section],1)
        if section=='entity': md(doc,model_section(notes,2))
        elif section=='boundary': md(doc,model_section(notes,3))
        else: md(doc,model_section(notes,4))
    doc.save(ROOT/'Deliverables.docx')

def build_ai():
    if not (ROOT/'ai-evidence/response.md').exists(): return
    doc=Document(); setup(doc)
    doc.add_heading('PlugPlan SG AI Technology Stack Critique',0)
    p(doc,'SC2006 Team SCE3-04 | Lab 2 | 11 September 2026')
    doc.add_heading('1 Material supplied and method',1)
    p(doc,'The AI received the Lab 1 system description with the agreed FR-24 and UC-07 moderation refinement, and the Lab 2 conceptual entity class diagram. The material covers Driver and Moderator roles, vehicle profiles, immutable Plan Versions, provider snapshots, recommendations and time-limited issue moderation. Full reproductions are retained in the repository rather than repeated here.')
    p(doc,'The response below was produced by an independent GPT-5.6-luna subagent configured at xhigh reasoning effort, using local file access.')
    doc.add_heading('2 Exact prompt',1); p(doc,(ROOT/'ai-evidence/prompt.txt').read_text(encoding='utf-8'))
    doc.add_page_break()
    doc.add_heading('3 Original AI response',1)
    p(doc,(ROOT/'ai-evidence/response.md').read_text(encoding='utf-8'))
    doc.add_heading('4 Opinion on one proposed technology',1)
    p(doc,(ROOT/'ai-evidence/opinion.txt').read_text(encoding='utf-8'))
    doc.save(ROOT/'AI-Critique-Report.docx')

if __name__=='__main__':
    import sys
    if '--ai-only' not in sys.argv: build_srs()
    build_ai()
